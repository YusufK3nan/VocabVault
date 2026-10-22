"""
Module: VocabVault ULTIMATE (Pomodoro & Export Edition)
Description: Premium Flashcard App with Custom Pomodoro, PDF/Word Export, Mastery, and Dynamic UI.
Dependencies: customtkinter, pywin32, reportlab, python-docx
"""

import os
import sys
import json
import random
import datetime
import threading
import winreg
import customtkinter as ctk
from tkinter import filedialog, messagebox

# PDF ve Word dışa aktarma kütüphaneleri kontrolü
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


# --- DİNAMİK TEMA RENKLERİ (Açık Mod, Koyu Mod) ---
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("green")

BG_COLOR = ("#F2F3F5", "#313338")        
SIDEBAR_BG = ("#E3E5E8", "#2B2D31")     
CARD_BG = ("#FFFFFF", "#1E1F22")        
ACCENT = ("#10B981", "#23A559")         
TEXT_COLOR = ("#000000", "#F2F3F5")
SUB_TEXT = ("#555555", "#A0A0A0")
HEADER_BG = ("#D1D5DB", "#3F3F46")


def get_desktop_path():
    try:
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders")
        desktop_path, _ = winreg.QueryValueEx(key, "Desktop")
        return os.path.expandvars(desktop_path)
    except Exception:
        return os.path.join(os.environ["USERPROFILE"], "Desktop")


def create_desktop_shortcut():
    try:
        if getattr(sys, 'frozen', False): exe_dir = os.path.dirname(sys.executable)
        else: exe_dir = os.path.dirname(os.path.abspath(__file__))

        config_path = os.path.join(exe_dir, "vocab_config.json")
        
        if os.path.exists(config_path):
            with open(config_path, "r", encoding="utf-8") as f:
                if json.load(f).get("shortcut_created_v2", False): return

        desktop_path = get_desktop_path()
        shortcut_path = os.path.join(desktop_path, "VocabVault.lnk")
        target = sys.executable

        import win32com.client
        wscript = win32com.client.Dispatch("WScript.Shell")
        shortcut = wscript.CreateShortCut(shortcut_path)
        shortcut.Targetpath = target
        shortcut.WorkingDirectory = os.path.dirname(target)
        shortcut.IconLocation = f"{target},0"
        shortcut.save()

        with open(config_path, "w", encoding="utf-8") as f:
            json.dump({"shortcut_created_v2": True}, f, indent=4)
    except Exception: pass


class VocabVaultApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("VocabVault Ultimate | Premium Learning")
        self.geometry("1050x700")
        self.resizable(False, False)
        self.configure(fg_color=BG_COLOR)
        
        if getattr(sys, 'frozen', False): self.exe_dir = os.path.dirname(sys.executable)
        else: self.exe_dir = os.path.dirname(os.path.abspath(__file__))
            
        self.data_file = os.path.join(self.exe_dir, "words_ultimate.json")
        self.leaderboard_file = os.path.join(self.exe_dir, "leaderboard.json")
        
        self.vocab_data = self.load_and_migrate_data()
        self.leaderboard_data = self.load_leaderboard()
        self.current_deck = "🌟 All Decks (Mix)"
        
        if not self.vocab_data:
            self.vocab_data = {"Genel": {}}
            self.save_data()

        # Quiz & Uygulama Durumları
        self.quiz_streak = 0
        self.quiz_active_word = ""
        self.quiz_active_deck = ""
        self.quiz_correct_answer = ""
        self.quiz_mode = "Standard"
        self.time_left = 60
        self.timer_running = False

        # Pomodoro Durumları
        self.pomo_time_left = 25 * 60
        self.pomo_running = False
        
        self.active_frame_name = ""
        self.flashcard_filter = "All"

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=1)

        self.deck_selector_var = ctk.StringVar(value=self.current_deck)

        self.setup_sidebar()
        
        self.frames = {}
        self.setup_flashcard_frame()
        self.setup_add_word_frame()
        self.setup_deck_manager_frame()
        self.setup_quiz_frame()
        
        self.select_frame("flashcard")
        self.bind_hotkeys()

    def bind_hotkeys(self):
        self.bind("<space>", self.hotkey_flip)
        self.bind("<Right>", self.hotkey_next)
        
    def hotkey_flip(self, event):
        if self.active_frame_name == "flashcard": self.flip_card()
            
    def hotkey_next(self, event):
        if self.active_frame_name == "flashcard": self.next_flashcard()

    # --- VERİ MİGRASYONU ---
    def load_and_migrate_data(self):
        old_data_file = os.path.join(self.exe_dir, "words_pro.json")
        data = {}
        if os.path.exists(old_data_file) and not os.path.exists(self.data_file):
            try:
                with open(old_data_file, "r", encoding="utf-8") as f:
                    old_data = json.load(f)
                for deck, words in old_data.items():
                    data[deck] = {}
                    for word, val in words.items():
                        if isinstance(val, str):
                            data[deck][word] = {"meaning": val, "example": "", "learned": False, "starred": False}
                        else:
                            data[deck][word] = val
                return data
            except: pass
            
        if os.path.exists(self.data_file):
            try:
                with open(self.data_file, "r", encoding="utf-8") as f:
                    return json.load(f)
            except: pass
        return {}

    def save_data(self):
        with open(self.data_file, "w", encoding="utf-8") as f:
            json.dump(self.vocab_data, f, indent=4, ensure_ascii=False)
            
    def load_leaderboard(self):
        if os.path.exists(self.leaderboard_file):
            try:
                with open(self.leaderboard_file, "r", encoding="utf-8") as f:
                    return json.load(f)
            except: pass
        return []

    def save_leaderboard(self):
        with open(self.leaderboard_file, "w", encoding="utf-8") as f:
            json.dump(self.leaderboard_data, f, indent=4, ensure_ascii=False)
            
    def add_to_leaderboard(self, score, mode):
        if score > 0:
            date_str = datetime.datetime.now().strftime("%d-%m %H:%M")
            self.leaderboard_data.append({"score": score, "date": date_str, "deck": self.current_deck, "mode": mode})
            self.leaderboard_data = sorted(self.leaderboard_data, key=lambda x: x["score"], reverse=True)[:15]
            self.save_leaderboard()

    # --- SİDEBAR ---
    def setup_sidebar(self):
        self.sidebar = ctk.CTkFrame(self, width=240, corner_radius=0, fg_color=SIDEBAR_BG)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_rowconfigure(7, weight=1)

        ctk.CTkLabel(self.sidebar, text="🧠 VocabVault", font=("Segoe UI", 26, "bold"), text_color=ACCENT).grid(row=0, column=0, padx=20, pady=(30, 35))

        btn_font = ("Segoe UI", 16, "bold")
        
        self.btn_menu_study = ctk.CTkButton(self.sidebar, text="🃏 Flashcards", height=45, font=btn_font, fg_color="transparent", hover_color=CARD_BG, text_color=TEXT_COLOR, anchor="w", command=lambda: self.select_frame("flashcard"))
        self.btn_menu_study.grid(row=1, column=0, padx=15, pady=6, sticky="ew")

        self.btn_menu_add = ctk.CTkButton(self.sidebar, text="✍️ Add Words", height=45, font=btn_font, fg_color="transparent", hover_color=CARD_BG, text_color=TEXT_COLOR, anchor="w", command=lambda: self.select_frame("add"))
        self.btn_menu_add.grid(row=2, column=0, padx=15, pady=6, sticky="ew")

        self.btn_menu_decks = ctk.CTkButton(self.sidebar, text="📁 Manage & Edit", height=45, font=btn_font, fg_color="transparent", hover_color=CARD_BG, text_color=TEXT_COLOR, anchor="w", command=lambda: self.select_frame("decks"))
        self.btn_menu_decks.grid(row=3, column=0, padx=15, pady=6, sticky="ew")
        
        self.btn_menu_quiz = ctk.CTkButton(self.sidebar, text="🔥 Quiz Arena", height=45, font=btn_font, fg_color=ACCENT, hover_color="#1e8b49", anchor="w", text_color="white", command=lambda: self.select_frame("quiz"))
        self.btn_menu_quiz.grid(row=4, column=0, padx=15, pady=20, sticky="ew")

        ctk.CTkLabel(self.sidebar, text="ACTIVE DECK", font=("Segoe UI", 12, "bold"), text_color=SUB_TEXT).grid(row=5, column=0, padx=20, pady=(0, 0), sticky="ws")
        
        self.deck_selector = ctk.CTkOptionMenu(self.sidebar, variable=self.deck_selector_var, values=self.get_deck_list_with_mix(), command=self.change_active_deck, fg_color=CARD_BG, button_color=HEADER_BG, text_color=TEXT_COLOR, font=("Segoe UI", 13))
        self.deck_selector.grid(row=6, column=0, padx=15, pady=(5, 10), sticky="ew")

        self.btn_theme = ctk.CTkButton(self.sidebar, text="🌙 / ☀️ Toggle Theme", height=35, text_color=TEXT_COLOR, fg_color="transparent", hover_color=CARD_BG, command=self.toggle_theme)
        self.btn_theme.grid(row=8, column=0, padx=15, pady=15, sticky="ew")

    def toggle_theme(self):
        current = ctk.get_appearance_mode()
        ctk.set_appearance_mode("Light" if current == "Dark" else "Dark")

    def get_deck_list_with_mix(self):
        return ["🌟 All Decks (Mix)"] + list(self.vocab_data.keys())

    def refresh_sidebar_decks(self):
        sidebar_decks = self.get_deck_list_with_mix()
        self.deck_selector.configure(values=sidebar_decks)
        
        normal_decks = list(self.vocab_data.keys())
        if hasattr(self, 'add_deck_selector'):
            self.add_deck_selector.configure(values=normal_decks)
            
        if self.current_deck not in sidebar_decks:
            self.current_deck = sidebar_decks[0] if sidebar_decks else "🌟 All Decks (Mix)"
            self.deck_selector_var.set(self.current_deck)

    def select_frame(self, frame_name):
        self.active_frame_name = frame_name
        for frame in self.frames.values(): frame.grid_remove()
        self.frames[frame_name].grid(row=0, column=1, sticky="nsew", padx=25, pady=25)
        
        if frame_name == "flashcard": self.next_flashcard()
        elif frame_name == "quiz":
            self.quiz_streak = 0
            self.timer_running = False
            self.load_quiz_question()
        elif frame_name == "decks": self.refresh_deck_list()

    def change_active_deck(self, new_deck):
        self.current_deck = new_deck
        self.deck_selector_var.set(new_deck)
        if self.active_frame_name == "flashcard": self.next_flashcard()
        elif self.active_frame_name == "quiz": self.load_quiz_question()

    # --- YARDIMCI METOTLAR ---
    def get_filtered_words(self, target_deck=None):
        deck_to_search = target_deck if target_deck else self.current_deck
        filtered = []
        if deck_to_search == "🌟 All Decks (Mix)":
            for d_name, d_words in self.vocab_data.items():
                for w, info in d_words.items():
                    if self.passes_filter(info): filtered.append((d_name, w, info))
        else:
            if deck_to_search in self.vocab_data:
                for w, info in self.vocab_data[deck_to_search].items():
                    if self.passes_filter(info): filtered.append((deck_to_search, w, info))
        return filtered

    def passes_filter(self, info):
        if self.flashcard_filter == "All": return True
        if self.flashcard_filter == "Starred": return info.get("starred", False)
        if self.flashcard_filter == "Unlearned": return not info.get("learned", False)
        return True

    def calculate_progress(self):
        words = self.get_filtered_words()
        if not words: return 0.0
        learned = sum(1 for _, _, info in words if info.get("learned", False))
        return learned / len(words)

    # --- 1. FLASHCARD (KARTLAR) EKRANI & POMODORO ---
    def setup_flashcard_frame(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        self.frames["flashcard"] = frame
        
        top_bar = ctk.CTkFrame(frame, fg_color="transparent")
        top_bar.pack(fill="x", pady=(0, 10))

        self.flashcard_title = ctk.CTkLabel(top_bar, text="Flashcards", text_color=TEXT_COLOR, font=("Segoe UI", 26, "bold"))
        self.flashcard_title.pack(side="left")

        self.filter_var = ctk.StringVar(value="All")
        ctk.CTkSegmentedButton(top_bar, values=["All", "Starred", "Unlearned"], variable=self.filter_var, command=self.change_filter, selected_color=ACCENT).pack(side="right")

        self.prog_bar = ctk.CTkProgressBar(frame, progress_color=ACCENT, height=8)
        self.prog_bar.pack(fill="x", pady=(0, 15), padx=20)
        self.prog_bar.set(0)

        self.card_bg = ctk.CTkFrame(frame, width=600, height=290, corner_radius=25, fg_color=CARD_BG, cursor="hand2")
        self.card_bg.pack(pady=5)
        self.card_bg.pack_propagate(False)

        self.lbl_card = ctk.CTkLabel(self.card_bg, text="", text_color=TEXT_COLOR, font=("Segoe UI", 42, "bold"), wraplength=550, cursor="hand2")
        self.lbl_card.place(relx=0.5, rely=0.42, anchor="center")
        
        self.lbl_example = ctk.CTkLabel(self.card_bg, text="", font=("Segoe UI", 16, "italic"), text_color=SUB_TEXT, wraplength=500)
        self.lbl_example.place(relx=0.5, rely=0.72, anchor="center")

        self.btn_star = ctk.CTkButton(self.card_bg, text="☆", width=40, height=40, font=("Arial", 24), fg_color="transparent", hover_color=HEADER_BG, text_color="#facc15", command=self.toggle_star)
        self.btn_star.place(relx=0.05, rely=0.05)
        
        self.btn_learned = ctk.CTkButton(self.card_bg, text="Tıkla: Öğrenildi İşaretle", width=120, height=30, fg_color="transparent", hover_color=HEADER_BG, text_color=SUB_TEXT, command=self.toggle_learned)
        self.btn_learned.place(relx=0.72, rely=0.05)

        self.card_bg.bind("<Button-1>", lambda e: self.flip_card())
        self.lbl_card.bind("<Button-1>", lambda e: self.flip_card())
        self.lbl_example.bind("<Button-1>", lambda e: self.flip_card())

        btn_panel = ctk.CTkFrame(frame, fg_color="transparent")
        btn_panel.pack(pady=15)
        
        ctk.CTkButton(btn_panel, text="🔊 Pronounce", width=140, height=50, font=("Segoe UI", 15, "bold"), fg_color="#3b82f6", hover_color="#2563eb", text_color="white", command=self.speak_word).grid(row=0, column=0, padx=15)
        ctk.CTkButton(btn_panel, text="Next Card ➡️", width=250, height=50, font=("Segoe UI", 18, "bold"), fg_color=ACCENT, hover_color="#1e8b49", text_color="white", command=self.next_flashcard).grid(row=0, column=1, padx=15)
        
        # --- ÖZELLEŞTİRİLEBİLİR POMODORO WIDGET (SAĞ ALT) ---
        pomo_frame = ctk.CTkFrame(frame, fg_color=CARD_BG, corner_radius=12)
        pomo_frame.pack(side="bottom", anchor="e", padx=10, pady=5)

        self.lbl_pomo = ctk.CTkLabel(pomo_frame, text="🍅 25:00", font=("Segoe UI", 16, "bold"), text_color=ACCENT)
        self.lbl_pomo.pack(side="left", padx=10, pady=5)

        self.pomo_min_menu = ctk.CTkOptionMenu(pomo_frame, values=["15 dk", "25 dk", "45 dk", "60 dk"], width=80, height=28, command=self.change_pomo_time)
        self.pomo_min_menu.set("25 dk")
        self.pomo_min_menu.pack(side="left", padx=5)

        self.btn_pomo_toggle = ctk.CTkButton(pomo_frame, text="Başlat", width=65, height=28, fg_color=ACCENT, hover_color="#1e8b49", text_color="white", command=self.toggle_pomodoro)
        self.btn_pomo_toggle.pack(side="left", padx=5)

        self.is_flipped = False
        self.card_active_word = None
        self.card_active_deck = None

    def change_pomo_time(self, choice):
        mins = int(choice.split()[0])
        self.pomo_time_left = mins * 60
        self.pomo_running = False
        self.btn_pomo_toggle.configure(text="Başlat")
        m, s = divmod(self.pomo_time_left, 60)
        self.lbl_pomo.configure(text=f"🍅 {m:02d}:{s:02d}")

    def toggle_pomodoro(self):
        if not self.pomo_running:
            self.pomo_running = True
            self.btn_pomo_toggle.configure(text="Durdur", fg_color="#dc3545")
            self.run_pomodoro_timer()
        else:
            self.pomo_running = False
            self.btn_pomo_toggle.configure(text="Başlat", fg_color=ACCENT)

    def run_pomodoro_timer(self):
        if not self.pomo_running: return
        if self.pomo_time_left > 0:
            self.pomo_time_left -= 1
            m, s = divmod(self.pomo_time_left, 60)
            self.lbl_pomo.configure(text=f"🍅 {m:02d}:{s:02d}")
            self.after(1000, self.run_pomodoro_timer)
        else:
            self.pomo_running = False
            self.btn_pomo_toggle.configure(text="Başlat", fg_color=ACCENT)
            messagebox.showinfo("Pomodoro Bitti!", "Harika odaklandın! Kısa bir mola vermeyi hak ettin ☕")

    def change_filter(self, new_filter):
        self.flashcard_filter = new_filter
        self.next_flashcard()

    def toggle_star(self):
        if self.card_active_word and self.card_active_deck:
            info = self.vocab_data[self.card_active_deck][self.card_active_word]
            info["starred"] = not info.get("starred", False)
            self.save_data()
            self.btn_star.configure(text="⭐" if info["starred"] else "☆")

    def toggle_learned(self):
        if self.card_active_word and self.card_active_deck:
            info = self.vocab_data[self.card_active_deck][self.card_active_word]
            info["learned"] = not info.get("learned", False)
            self.save_data()
            self.btn_learned.configure(text="✅ Learned" if info["learned"] else "⭕ Unlearned", text_color=ACCENT if info["learned"] else SUB_TEXT)
            self.prog_bar.set(self.calculate_progress())

    def speak_word(self):
        if self.card_active_word and not self.is_flipped:
            word = self.card_active_word
            def run_tts():
                try:
                    import win32com.client
                    speaker = win32com.client.Dispatch("SAPI.SpVoice")
                    speaker.Speak(word)
                except: pass
            threading.Thread(target=run_tts, daemon=True).start()

    def flip_card(self):
        if not self.card_active_word: return
        info = self.vocab_data[self.card_active_deck][self.card_active_word]

        if not self.is_flipped:
            self.card_bg.configure(fg_color=("#dcfce7", "#14532d")) 
            self.lbl_card.configure(text=info["meaning"].capitalize(), text_color=("#166534", "#4ade80"))
            self.lbl_example.configure(text="")
            self.is_flipped = True
        else:
            self.card_bg.configure(fg_color=CARD_BG) 
            self.lbl_card.configure(text=self.card_active_word.capitalize(), text_color=TEXT_COLOR)
            ex = info.get("example", "")
            self.lbl_example.configure(text=f'"{ex}"' if ex else "")
            self.is_flipped = False

    def next_flashcard(self):
        words_pool = self.get_filtered_words()
        self.prog_bar.set(self.calculate_progress())
        self.card_bg.configure(fg_color=CARD_BG) 
        
        if not words_pool:
            self.lbl_card.configure(text="No words found matching criteria.", text_color=SUB_TEXT, font=("Segoe UI", 24))
            self.lbl_example.configure(text="")
            self.card_active_word = None
            return

        deck, word, info = random.choice(words_pool)
        self.card_active_deck, self.card_active_word = deck, word
        
        self.lbl_card.configure(text=word.capitalize(), text_color=TEXT_COLOR, font=("Segoe UI", 42, "bold"))
        ex = info.get("example", "")
        self.lbl_example.configure(text=f'"{ex}"' if ex else "")
        self.btn_star.configure(text="⭐" if info.get("starred") else "☆")
        self.btn_learned.configure(text="✅ Learned" if info.get("learned") else "⭕ Unlearned", text_color=ACCENT if info.get("learned") else SUB_TEXT)
        self.is_flipped = False

    # --- 2. KELİME EKLEME EKRANI ---
    def setup_add_word_frame(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        self.frames["add"] = frame

        ctk.CTkLabel(frame, text="Add New Vocabulary", text_color=TEXT_COLOR, font=("Segoe UI", 28, "bold")).pack(pady=(20, 10))
        
        deck_frame = ctk.CTkFrame(frame, fg_color="transparent")
        deck_frame.pack(pady=10)
        ctk.CTkLabel(deck_frame, text="Target Deck:", text_color=TEXT_COLOR, font=("Segoe UI", 15)).pack(side="left", padx=10)
        
        self.add_deck_selector = ctk.CTkOptionMenu(deck_frame, variable=self.deck_selector_var, values=list(self.vocab_data.keys()), fg_color=CARD_BG, button_color=HEADER_BG, text_color=TEXT_COLOR, width=250)
        self.add_deck_selector.pack(side="left")

        self.entry_eng = ctk.CTkEntry(frame, placeholder_text="English Word (e.g. Obsolete)", fg_color=CARD_BG, text_color=TEXT_COLOR, width=500, height=55, font=("Segoe UI", 16))
        self.entry_eng.pack(pady=10)

        self.entry_turk = ctk.CTkEntry(frame, placeholder_text="Turkish Meaning (e.g. Modası geçmiş)", fg_color=CARD_BG, text_color=TEXT_COLOR, width=500, height=55, font=("Segoe UI", 16))
        self.entry_turk.pack(pady=10)
        
        self.entry_ex = ctk.CTkEntry(frame, placeholder_text="Example Sentence (Optional)", fg_color=CARD_BG, text_color=TEXT_COLOR, width=500, height=55, font=("Segoe UI", 16))
        self.entry_ex.pack(pady=10)

        ctk.CTkButton(frame, text="💾 Save to Vault", width=500, height=60, font=("Segoe UI", 18, "bold"), fg_color=ACCENT, hover_color="#1e8b49", text_color="white", command=self.add_word).pack(pady=25)
        self.lbl_add_status = ctk.CTkLabel(frame, text="", font=("Segoe UI", 15))
        self.lbl_add_status.pack()

    def add_word(self):
        target = self.deck_selector_var.get()
        if target == "🌟 All Decks (Mix)":
            self.lbl_add_status.configure(text="❌ Please select a specific deck!", text_color="#ff4c4c")
            return

        eng = self.entry_eng.get().strip().lower()
        turk = self.entry_turk.get().strip().lower()
        ex = self.entry_ex.get().strip()

        if not eng or not turk:
            self.lbl_add_status.configure(text="❌ Word and Meaning are required!", text_color="#ff4c4c")
            return

        self.vocab_data[target][eng] = {"meaning": turk, "example": ex, "learned": False, "starred": False}
        self.save_data()

        self.entry_eng.delete(0, "end")
        self.entry_turk.delete(0, "end")
        self.entry_ex.delete(0, "end")
        self.entry_eng.focus()
        
        self.lbl_add_status.configure(text=f"✅ '{eng.capitalize()}' added!", text_color=ACCENT)

    # --- 3. KLASÖR YÖNETİMİ & ÇİFT FORMATLI ÇIKTI (PDF & WORD) ---
    def setup_deck_manager_frame(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        self.frames["decks"] = frame
        
        top_bar = ctk.CTkFrame(frame, fg_color="transparent")
        top_bar.pack(fill="x", pady=10)
        ctk.CTkLabel(top_bar, text="Manage & Edit", text_color=TEXT_COLOR, font=("Segoe UI", 26, "bold")).pack(side="left")
        
        # Çıktı ve Dosya Butonları
        ctk.CTkButton(top_bar, text="📝 Word", width=80, fg_color="#1d4ed8", text_color="white", command=self.export_to_word).pack(side="right", padx=3)
        ctk.CTkButton(top_bar, text="📄 PDF", width=80, fg_color="#dc2626", text_color="white", command=self.export_to_pdf).pack(side="right", padx=3)
        ctk.CTkButton(top_bar, text="📤 Export", width=80, fg_color="#3b82f6", text_color="white", command=self.export_deck).pack(side="right", padx=3)
        ctk.CTkButton(top_bar, text="📥 Import", width=80, fg_color="#8b5cf6", text_color="white", command=self.import_deck).pack(side="right", padx=3)

        tools = ctk.CTkFrame(frame, fg_color="transparent")
        tools.pack(fill="x", pady=5)
        
        self.entry_new_deck = ctk.CTkEntry(tools, placeholder_text="New Deck Name", fg_color=CARD_BG, text_color=TEXT_COLOR, width=180)
        self.entry_new_deck.pack(side="left", padx=5)
        ctk.CTkButton(tools, text="Create", width=70, fg_color=HEADER_BG, text_color=TEXT_COLOR, command=self.create_deck).pack(side="left")
        
        self.entry_search = ctk.CTkEntry(tools, placeholder_text="🔍 Search words...", fg_color=CARD_BG, text_color=TEXT_COLOR, width=200)
        self.entry_search.pack(side="right", padx=5)
        self.entry_search.bind("<KeyRelease>", lambda e: self.refresh_deck_list(self.entry_search.get().lower()))
        
        self.deck_list_frame = ctk.CTkScrollableFrame(frame, width=680, height=400, fg_color=CARD_BG)
        self.deck_list_frame.pack(pady=15)

    def export_to_pdf(self):
        if not PDF_AVAILABLE:
            messagebox.showerror("Eksik Kütüphane", "PDF çıktısı için 'reportlab' kurulu değil!\nTerminalden 'pip install reportlab' çalıştırın.")
            return
        
        target = self.deck_selector_var.get()
        if target == "🌟 All Decks (Mix)":
            words_to_print = []
            for d, w_dict in self.vocab_data.items():
                for w, info in w_dict.items(): words_to_print.append((d, w, info))
        else:
            if target not in self.vocab_data or not self.vocab_data[target]:
                messagebox.showwarning("Boş Deste", "Seçili deste boş veya seçilmedi!")
                return
            words_to_print = [(target, w, info) for w, info in self.vocab_data[target].items()]

        filepath = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")])
        if filepath:
            try:
                c = canvas.Canvas(filepath, pagesize=letter)
                width, height = letter
                
                c.setFont("Helvetica-Bold", 20)
                c.drawString(50, height - 50, f"VocabVault - Study Sheet ({target})")
                
                c.setFont("Helvetica", 12)
                y = height - 90
                for d_name, w, info in words_to_print:
                    if y < 50:
                        c.showPage()
                        c.setFont("Helvetica", 12)
                        y = height - 50
                    text_line = f"• {w.capitalize()} : {info['meaning'].capitalize()}"
                    if info.get('example'): text_line += f"   (Ex: {info['example']})"
                    c.drawString(50, y, text_line)
                    y -= 25
                c.save()
                messagebox.showinfo("Başarılı", "PDF çalışma kağıdı başarıyla oluşturuldu!")
            except Exception as e:
                messagebox.showerror("Hata", f"PDF oluşturulamadı: {str(e)}")

    def export_to_word(self):
        if not DOCX_AVAILABLE:
            messagebox.showerror("Eksik Kütüphane", "Word çıktısı için 'python-docx' kurulu değil!\nTerminalden 'pip install python-docx' çalıştırın.")
            return

        target = self.deck_selector_var.get()
        if target == "🌟 All Decks (Mix)":
            words_to_print = []
            for d, w_dict in self.vocab_data.items():
                for w, info in w_dict.items(): words_to_print.append((d, w, info))
        else:
            if target not in self.vocab_data or not self.vocab_data[target]:
                messagebox.showwarning("Boş Deste", "Seçili deste boş veya seçilmedi!")
                return
            words_to_print = [(target, w, info) for w, info in self.vocab_data[target].items()]

        filepath = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filepath:
            try:
                doc = Document()
                doc.add_heading(f"VocabVault - Study Sheet ({target})", level=1)
                
                for d_name, w, info in words_to_print:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(f"{w.capitalize()}").bold = True
                    p.add_run(f" - {info['meaning'].capitalize()}")
                    if info.get('example'):
                        p.add_run(f"\n   Örnek: {info['example']}").italic = True
                
                doc.save(filepath)
                messagebox.showinfo("Başarılı", "Word çalışma kağıdı başarıyla oluşturuldu!")
            except Exception as e:
                messagebox.showerror("Hata", f"Word dosyası oluşturulamadı: {str(e)}")

    def import_deck(self):
        filepath = filedialog.askopenfilename(filetypes=[("JSON Files", "*.json")])
        if filepath:
            try:
                with open(filepath, "r", encoding="utf-8") as f:
                    new_decks = json.load(f)
                for d_name, d_words in new_decks.items():
                    if d_name not in self.vocab_data: self.vocab_data[d_name] = {}
                    for w, val in d_words.items():
                        if isinstance(val, str): self.vocab_data[d_name][w] = {"meaning": val, "example": "", "learned": False, "starred": False}
                        else: self.vocab_data[d_name][w] = val
                self.save_data()
                self.refresh_sidebar_decks()
                self.refresh_deck_list()
                messagebox.showinfo("Success", "Decks imported successfully!")
            except: messagebox.showerror("Error", "Invalid file format.")

    def export_deck(self):
        target = self.deck_selector_var.get()
        if target == "🌟 All Decks (Mix)":
            export_data = self.vocab_data
        else:
            export_data = {target: self.vocab_data[target]}
            
        filepath = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON Files", "*.json")])
        if filepath:
            with open(filepath, "w", encoding="utf-8") as f:
                json.dump(export_data, f, indent=4, ensure_ascii=False)
            messagebox.showinfo("Success", "Deck exported successfully!")

    def create_deck(self):
        new_deck = self.entry_new_deck.get().strip()
        if new_deck and new_deck not in self.vocab_data:
            self.vocab_data[new_deck] = {}
            self.save_data()
            self.entry_new_deck.delete(0, "end")
            self.refresh_sidebar_decks()
            self.refresh_deck_list()

    def refresh_deck_list(self, search_query=""):
        for widget in self.deck_list_frame.winfo_children(): widget.destroy()
            
        for deck_name, words in self.vocab_data.items():
            filtered_words = {k: v for k, v in words.items() if search_query in k or search_query in v["meaning"]} if search_query else words
            if search_query and not filtered_words: continue

            header_row = ctk.CTkFrame(self.deck_list_frame, fg_color=HEADER_BG, corner_radius=8)
            header_row.pack(fill="x", pady=(10, 2), padx=5)
            
            ctk.CTkLabel(header_row, text=f"📁 {deck_name.upper()} ({len(filtered_words)})", font=("Segoe UI", 16, "bold"), text_color=ACCENT).pack(side="left", padx=15, pady=10)
            if deck_name != "Genel":
                ctk.CTkButton(header_row, text="Del Deck", width=70, fg_color="#dc3545", hover_color="#c82333", text_color="white", command=lambda d=deck_name: self.delete_deck(d)).pack(side="right", padx=5)

            content_frame = ctk.CTkFrame(self.deck_list_frame, fg_color="transparent")
            btn_toggle = ctk.CTkButton(header_row, text="🔽", width=35, fg_color="transparent", hover_color=SIDEBAR_BG, text_color=TEXT_COLOR)
            btn_toggle.configure(command=lambda c=content_frame, b=btn_toggle: self.toggle_accordion(c, b))
            btn_toggle.pack(side="right", padx=5)
            
            if search_query: content_frame.pack(fill="x", padx=10, pady=2)

            for eng, info in filtered_words.items():
                word_row = ctk.CTkFrame(content_frame, fg_color="transparent")
                word_row.pack(fill="x", pady=2, padx=15)
                
                status = "⭐ " if info.get("starred") else ""
                status += "✅ " if info.get("learned") else ""
                
                ctk.CTkLabel(word_row, text=f"{status}{eng.capitalize()} - {info['meaning'].capitalize()}", text_color=TEXT_COLOR, font=("Segoe UI", 14)).pack(side="left")
                
                ctk.CTkButton(word_row, text="❌", width=28, fg_color="transparent", text_color="#dc3545", command=lambda d=deck_name, w=eng: self.delete_word(d, w)).pack(side="right")
                ctk.CTkButton(word_row, text="✏️ Edit", width=50, fg_color="transparent", text_color="#3b82f6", command=lambda d=deck_name, w=eng, i=info: self.edit_word_popup(d, w, i)).pack(side="right", padx=10)

    def edit_word_popup(self, deck_name, old_word, info):
        win = ctk.CTkToplevel(self)
        win.title("Edit Word")
        win.geometry("400x350")
        win.grab_set()
        win.configure(fg_color=BG_COLOR)

        ctk.CTkLabel(win, text="English Word:", text_color=TEXT_COLOR).pack(pady=(15, 0))
        entry_w = ctk.CTkEntry(win, width=300, fg_color=CARD_BG, text_color=TEXT_COLOR)
        entry_w.pack(); entry_w.insert(0, old_word)
        
        ctk.CTkLabel(win, text="Turkish Meaning:", text_color=TEXT_COLOR).pack(pady=(10, 0))
        entry_m = ctk.CTkEntry(win, width=300, fg_color=CARD_BG, text_color=TEXT_COLOR)
        entry_m.pack(); entry_m.insert(0, info["meaning"])
        
        ctk.CTkLabel(win, text="Example:", text_color=TEXT_COLOR).pack(pady=(10, 0))
        entry_ex = ctk.CTkEntry(win, width=300, fg_color=CARD_BG, text_color=TEXT_COLOR)
        entry_ex.pack(); entry_ex.insert(0, info.get("example", ""))
        
        def save_edit():
            new_w = entry_w.get().strip().lower()
            if new_w and new_w != old_word:
                self.vocab_data[deck_name][new_w] = self.vocab_data[deck_name].pop(old_word)
            else:
                new_w = old_word
            
            self.vocab_data[deck_name][new_w]["meaning"] = entry_m.get().strip().lower()
            self.vocab_data[deck_name][new_w]["example"] = entry_ex.get().strip()
            self.save_data()
            self.refresh_deck_list()
            win.destroy()
            
        ctk.CTkButton(win, text="Update", fg_color=ACCENT, text_color="white", command=save_edit).pack(pady=25)

    def toggle_accordion(self, content_frame, btn):
        if content_frame.winfo_ismapped():
            content_frame.pack_forget(); btn.configure(text="🔽")
        else:
            content_frame.pack(fill="x", padx=10, pady=2); btn.configure(text="🔼")

    def delete_deck(self, d):
        if messagebox.askyesno("Delete", f"Delete entire deck '{d}'?"):
            if d in self.vocab_data:
                del self.vocab_data[d]
                self.save_data()
                self.refresh_sidebar_decks()
                self.refresh_deck_list()
            
    def delete_word(self, d, w):
        if d in self.vocab_data and w in self.vocab_data[d]:
            del self.vocab_data[d][w]
            self.save_data()
            self.refresh_deck_list(self.entry_search.get().lower())

    # --- 4. QUIZ ARENA ---
    def setup_quiz_frame(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        self.frames["quiz"] = frame

        top_bar = ctk.CTkFrame(frame, fg_color="transparent")
        top_bar.pack(fill="x", pady=10)
        
        self.quiz_mode_var = ctk.StringVar(value="Standard")
        modes = ctk.CTkSegmentedButton(top_bar, values=["Standard", "Type-in", "60s Rush"], variable=self.quiz_mode_var, command=self.change_quiz_mode, selected_color=ACCENT)
        modes.pack(side="left")
        
        self.lbl_streak_timer = ctk.CTkLabel(top_bar, text="🔥 Streak: 0", font=("Segoe UI", 20, "bold"), text_color="#f97316")
        self.lbl_streak_timer.pack(side="right", padx=15)

        self.quiz_play_area = ctk.CTkFrame(frame, fg_color=CARD_BG, corner_radius=20, height=350)
        self.quiz_play_area.pack(fill="x", pady=20, padx=10)
        self.quiz_play_area.pack_propagate(False)
        
        self.lbl_quiz_word = ctk.CTkLabel(self.quiz_play_area, text="", text_color=TEXT_COLOR, font=("Segoe UI", 36, "bold"))
        self.lbl_quiz_word.pack(pady=30)
        
        self.quiz_buttons_frame = ctk.CTkFrame(self.quiz_play_area, fg_color="transparent")
        self.quiz_buttons = []
        for i in range(4):
            btn = ctk.CTkButton(self.quiz_buttons_frame, text="", width=250, height=55, font=("Segoe UI", 16, "bold"), text_color="white", command=lambda idx=i: self.check_quiz_answer(idx))
            btn.grid(row=i//2, column=i%2, padx=15, pady=10)
            self.quiz_buttons.append(btn)
            
        self.type_in_frame = ctk.CTkFrame(self.quiz_play_area, fg_color="transparent")
        self.entry_type_in = ctk.CTkEntry(self.type_in_frame, width=300, height=50, font=("Segoe UI", 18), fg_color=BG_COLOR, text_color=TEXT_COLOR, placeholder_text="Type English word...")
        self.entry_type_in.pack(pady=10)
        self.entry_type_in.bind("<Return>", lambda e: self.check_type_in_answer())
            
        self.lbl_quiz_feedback = ctk.CTkLabel(frame, text="", font=("Segoe UI", 18, "bold"))
        self.lbl_quiz_feedback.pack(pady=5)
        
        ctk.CTkButton(frame, text="🏆 Leaderboards", width=220, height=45, font=("Segoe UI", 15, "bold"), fg_color="#eab308", text_color="black", hover_color="#ca8a04", command=self.show_leaderboard).pack(pady=10)

    def change_quiz_mode(self, mode):
        self.quiz_mode = mode
        self.quiz_streak = 0
        self.timer_running = False
        self.time_left = 60
        self.load_quiz_question()

    def update_timer(self):
        if not self.timer_running: return
        if self.time_left > 0:
            self.time_left -= 1
            self.lbl_streak_timer.configure(text=f"⏱️ {self.time_left}s | 🔥 {self.quiz_streak}", text_color="#ef4444" if self.time_left <= 10 else "#f97316")
            self.after(1000, self.update_timer)
        else:
            self.timer_running = False
            self.add_to_leaderboard(self.quiz_streak, "Rush")
            self.lbl_quiz_word.configure(text="TIME'S UP!", text_color="#ef4444")
            self.quiz_buttons_frame.pack_forget()
            self.lbl_quiz_feedback.configure(text=f"Final Score: {self.quiz_streak}", text_color=ACCENT)

    def load_quiz_question(self):
        words_pool = self.get_filtered_words()
        
        if self.quiz_mode == "60s Rush":
            if not self.timer_running and self.time_left == 60:
                self.timer_running = True
                self.update_timer()
            self.lbl_streak_timer.configure(text=f"⏱️ {self.time_left}s | 🔥 {self.quiz_streak}")
        else:
            self.lbl_streak_timer.configure(text=f"🔥 Streak: {self.quiz_streak}", text_color="#f97316")
        
        if len(words_pool) < 4:
            self.lbl_quiz_word.configure(text="⚠️ Not enough words!", text_color="#facc15")
            self.quiz_buttons_frame.pack_forget()
            self.type_in_frame.pack_forget()
            return
            
        self.lbl_quiz_feedback.configure(text="")
        
        if self.quiz_mode == "Type-in":
            self.quiz_buttons_frame.pack_forget()
            self.type_in_frame.pack()
            self.entry_type_in.delete(0, "end")
            self.entry_type_in.focus()
            
            deck, eng, info = random.choice(words_pool)
            self.quiz_active_word = eng
            self.quiz_correct_answer = eng
            self.lbl_quiz_word.configure(text=f"Türkçesi: {info['meaning'].capitalize()}", text_color=TEXT_COLOR)
        else:
            self.type_in_frame.pack_forget()
            self.quiz_buttons_frame.pack()
            for btn in self.quiz_buttons: btn.configure(state="normal", fg_color=["#3b82f6", "#1f538d"])
            
            deck, eng, info = random.choice(words_pool)
            self.quiz_active_word = eng
            self.quiz_correct_answer = info["meaning"]
            self.lbl_quiz_word.configure(text=eng.capitalize(), text_color=TEXT_COLOR)
            
            wrongs = [i["meaning"] for d, w, i in words_pool if i["meaning"] != self.quiz_correct_answer]
            if len(wrongs) < 3: wrongs += ["N/A", "Unknown", "Empty"]
            options = random.sample(wrongs, 3) + [self.quiz_correct_answer]
            random.shuffle(options)
            
            for i in range(4): self.quiz_buttons[i].configure(text=options[i].capitalize())

    def check_type_in_answer(self):
        ans = self.entry_type_in.get().strip().lower()
        if ans == self.quiz_correct_answer:
            self.quiz_streak += 1
            self.lbl_quiz_feedback.configure(text="✅ Perfect!", text_color=ACCENT)
            self.after(500, self.load_quiz_question)
        else:
            self.lbl_quiz_feedback.configure(text=f"❌ Wrong! Correct: {self.quiz_correct_answer.capitalize()}", text_color="#ef4444")
            self.add_to_leaderboard(self.quiz_streak, "Type-in")
            self.quiz_streak = 0
            self.after(2000, self.load_quiz_question)

    def check_quiz_answer(self, btn_index):
        if self.quiz_mode == "60s Rush" and not self.timer_running: return
        
        ans = self.quiz_buttons[btn_index].cget("text").lower()
        if ans == self.quiz_correct_answer.lower():
            self.quiz_buttons[btn_index].configure(fg_color=ACCENT) 
            self.quiz_streak += 1
            self.lbl_quiz_feedback.configure(text="✅ Correct!", text_color=ACCENT)
            self.after(400, self.load_quiz_question) 
        else:
            self.quiz_buttons[btn_index].configure(fg_color="#ef4444")
            self.lbl_quiz_feedback.configure(text=f"❌ Correct was: {self.quiz_correct_answer.capitalize()}", text_color="#ef4444")
            if self.quiz_mode != "60s Rush":
                self.add_to_leaderboard(self.quiz_streak, "Standard")
                self.quiz_streak = 0
            self.after(1500, self.load_quiz_question)

    def show_leaderboard(self):
        win = ctk.CTkToplevel(self)
        win.title("Leaderboards")
        win.geometry("500x550")
        win.grab_set()
        win.configure(fg_color=BG_COLOR)
        
        ctk.CTkLabel(win, text="🏆 Hall of Fame", font=("Segoe UI", 26, "bold"), text_color="#facc15").pack(pady=15)
        scroll = ctk.CTkScrollableFrame(win, width=420, height=400, fg_color=CARD_BG)
        scroll.pack(pady=10)
        
        if not self.leaderboard_data: return
            
        for i, rec in enumerate(self.leaderboard_data):
            c = "#facc15" if i==0 else "#9ca3af" if i==1 else "#b45309" if i==2 else TEXT_COLOR
            row = ctk.CTkFrame(scroll, fg_color="transparent")
            row.pack(fill="x", pady=5)
            ctk.CTkLabel(row, text=f"#{i+1}", font=("Segoe UI", 18, "bold"), text_color=c).pack(side="left", padx=10)
            ctk.CTkLabel(row, text=f"{rec['score']} pts", font=("Segoe UI", 16, "bold"), text_color=ACCENT).pack(side="left")
            ctk.CTkLabel(row, text=f"[{rec['mode']}] {rec['deck'][:15]}", font=("Segoe UI", 12), text_color=SUB_TEXT).pack(side="right", padx=10)

if __name__ == "__main__":
    create_desktop_shortcut()
    app = VocabVaultApp()
    app.mainloop()